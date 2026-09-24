# -*- coding: utf-8 -*-
import csv, json, re, os
from collections import defaultdict
from copy import deepcopy

SRC = "/home/workdir/attachments/tamdad_questions(1).csv"
OUT = "/home/workdir/artifacts/vakalat-quiz/data"

# Manual grouping: primary article + title
# Based on content, not just first mentioned article
MANUAL = {
"29278": ("ماده ۱۹ ق.آ.د.م", "قرار اناطه", "آسان", "اناطه-19"),
"3236592": ("ماده ۱۹ ق.آ.د.م", "قرار اناطه", "آسان", "اناطه-19"),
"29279": ("ماده ۲۰ ق.آ.د.م", "صلاحیت دعاوی ترکه", "آسان", "ترکه-20"),
"29280": ("ماده ۲۰ ق.آ.د.م", "صلاحیت دعاوی ترکه", "سخت", "ترکه-20"),
"29281": ("ماده ۲۰ ق.آ.د.م", "صلاحیت دعاوی ترکه", "متوسط", "ترکه-20"),
"29282": ("ماده ۲۰ ق.آ.د.م", "صلاحیت دعاوی ترکه", "آسان", "ترکه-20"),
"29283": ("ماده ۲۰ ق.آ.د.م", "صلاحیت دعاوی ترکه", "سخت", "ترکه-20"),
"29284": ("مواد ۴۷۵ و ۵۸۳ ق.م", "افراز ملک اجاره‌ای", "سخت", "افراز-اجاره"),
"29285": ("ماده ۲۱ ق.آ.د.م", "صلاحیت دعوای ورشکستگی", "آسان", "ورشکستگی-21"),
"29286": ("ماده ۲۱ ق.آ.د.م", "صلاحیت دعوای ورشکستگی", "آسان", "ورشکستگی-21"),
"29305": ("مواد ۱۱ و ۲۱ ق.آ.د.م", "ورشکستگی شریک مقیم خارج", "سخت", "ورشکستگی-21"),
"29306": ("مواد ۱۱ و ۲۱ ق.آ.د.م", "ورشکستگی تاجر مقیم خارج", "سخت", "ورشکستگی-21"),
"29303": ("ماده ۲۱ ق.آ.د.م", "ورشکستگی شریک شرکت", "متوسط", "ورشکستگی-21"),
"29287": ("ماده ۲۲ ق.آ.د.م", "اختلاف شرکا در شرکت", "آسان", "شرکت-22-23"),
"29288": ("ماده ۲۳ ق.آ.د.م", "تعهدات شرکت در برابر اشخاص ثالث", "متوسط", "شرکت-22-23"),
"29289": ("ماده ۵۰۵ ق.آ.د.م", "اعسار از هزینه تجدیدنظر", "آسان", "اعسار-24-505"),
"29290": ("ماده ۲۴ ق.آ.د.م", "اعسار از محکوم‌به", "آسان", "اعسار-24-505"),
"29291": ("ماده ۲۴ ق.آ.د.م", "اعسار از هزینه تجدیدنظر", "آسان", "اعسار-24-505"),
"29292": ("ماده ۵۰۵ ق.آ.د.م", "اعسار از هزینه تجدیدنظر", "آسان", "اعسار-24-505"),
"3236815": ("ماده ۲۴ ق.آ.د.م", "مرجع دعوای اعسار", "آسان", "اعسار-24-505"),
"29293": ("ماده ۴ ق.ثبت احوال + بند ۱۱ م.۱۲ شورا", "تغییر نام / اسناد سجلی", "سخت", "ثبت-احوال-25"),
"29297": ("ماده ۲۵ ق.آ.د.م", "اسناد سجلی ذی‌نفع مقیم خارج", "آسان", "ثبت-احوال-25"),
"29298": ("ماده ۲۵ ق.آ.د.م", "اسناد سجلی ذی‌نفع مقیم خارج", "آسان", "ثبت-احوال-25"),
"29299": ("ماده ۲۵ ق.آ.د.م + دادگاه صلح", "سند سجلی تنظیم‌شده در خارج", "متوسط", "ثبت-احوال-25"),
"29294": ("ماده ۲۵ ق.آ.د.م + دادگاه صلح", "سند سجلی تنظیم‌شده در خارج", "متوسط", "ثبت-احوال-25"),
"29296": ("ماده ۲۵ ق.آ.د.م", "اسناد سجلی ذی‌نفع مقیم خارج", "آسان", "ثبت-احوال-25"),
"29302": ("اصل ۱۵۹ + ماده ۱ افراز", "افراز ملک قولنامه‌ای", "متوسط", "افراز-قولنامه"),
"29301": ("اصل ۱۵۹ + ماده ۱ افراز", "افراز ملک قولنامه‌ای", "متوسط", "افراز-قولنامه"),
"29307": ("مواد ۱۱ و ۱۶ ق.آ.د.م", "خواندگان متعدد / خسارت کارگران", "سخت", "صلاحیت-عمومی-11-16"),
"29308": ("اصل ۱۵۹ + م.۱۰ دیوان", "دعوای قراردادی با دستگاه دولتی", "سخت", "دیوان-قرارداد"),
"29309": ("م.۴ حمایت خانواده + م.۱۰ دیوان + م.۱۲ شورا", "صلاحیت ترکیبی خانواده/دیوان/صلح", "سخت", "صلاحیت-ترکیبی"),
"29310": ("ماده ۱۲ قانون دیوان عدالت اداری", "اعتراض به آیین‌نامه دولتی", "متوسط", "دیوان-آیین‌نامه"),
"3219937": ("مواد ۱۰ و ۳۴ قانون دیوان", "شکایت استخدامی و دستور موقت دیوان", "سخت", "دیوان-دستور-موقت"),
"3219938": ("مواد ۱۵ و ۲۱ قانون شوراهای حل اختلاف ۱۴۰۲", "ارجاع رجوع در طلاق به شورا", "سخت", "شورا-خانواده"),
"29256": ("مواد ۱۱ و ۱۳ ق.آ.د.م", "استرداد مال منقول ناشی از قرارداد", "آسان", "قرارداد-منقول-13"),
"29258": ("مواد ۱۱ و ۱۳ ق.آ.د.م", "صلاحیت دعوای چک", "سخت", "قرارداد-منقول-13"),
"29304": ("مواد ۱۱ و ۱۳ ق.آ.د.م", "خواندگان متعدد + قرارداد منقول", "سخت", "قرارداد-منقول-13"),
"29271": ("مواد ۱۷ و ۱۴۱ ق.آ.د.م", "دعوای اضافی / خسارت مورد اجاره", "سخت", "طاری-17"),
"29274": ("ماده ۱۷ ق.آ.د.م", "اقامه دعوای طاری", "متوسط", "طاری-17"),
"29267": ("ماده ۱۶ ق.آ.د.م", "اموال غیرمنقول متعدد", "آسان", "صلاحیت-16"),
"29268": ("ماده ۱۶ ق.آ.د.م", "خواندگان متعدد", "آسان", "صلاحیت-16"),
"29260": ("ماده ۱۵ ق.آ.د.م", "منقول و غیرمنقول ناشی از یک منشأ", "متوسط", "صلاحیت-15"),
"29264": ("ماده ۱۵ ق.آ.د.م", "منقول و غیرمنقول ناشی از یک منشأ", "آسان", "صلاحیت-15"),
"29266": ("ماده ۱۵ ق.آ.د.م + ماده ۲۰ ق.م", "اجرت‌المسمی و اجرت‌المثل", "سخت", "صلاحیت-15"),
"29269": ("ماده ۱۷ ق.آ.د.م", "انواع دعاوی طاری", "آسان", "طاری-17"),
"29270": ("ماده ۱۷ ق.آ.د.م", "تعریف دعوای طاری", "آسان", "طاری-17"),
"29272": ("مواد ۱۷ و ۱۴۱ ق.آ.د.م", "مرجع رسیدگی به دعوای طاری", "متوسط", "طاری-17"),
"29276": ("ماده ۱۷ ق.آ.د.م", "انواع دعاوی طاری", "آسان", "طاری-17"),
"29277": ("ماده ۱۷ ق.آ.د.م", "دعوای اضافی", "آسان", "طاری-17"),
}

DIFF_ORDER = {"سخت": 0, "متوسط": 1, "آسان": 2}

with open(SRC, encoding="utf-8-sig") as f:
    rows = list(csv.DictReader(f))

questions = []
for row in rows:
    qid = str(row["ID"]).strip()
    art, title, diff, group = MANUAL.get(qid, ("سایر", "سایر", "متوسط", "other"))
    questions.append({
        "id": qid,
        "row": row["ردیف"],
        "lesson": "آیین دادرسی مدنی",
        "lesson_id": "adm",
        "question": row["سوال"].strip(),
        "options": {
            "A": row["گزینه A"].strip(),
            "B": row["گزینه B"].strip(),
            "C": row["گزینه C"].strip(),
            "D": row["گزینه D"].strip(),
        },
        "answer_text": row["پاسخ صحیح"].strip(),
        "answer": row["Choice صحیح"].replace("Choice", "").strip(),
        "explain": row["پاسخ تشریحی"].strip(),
        "laws": row["قوانین مرتبط"].strip(),
        "article": art,
        "article_title": title,
        "difficulty": diff,
        "similar_group": group,
    })

# mark similar
groups = defaultdict(list)
for q in questions:
    groups[q["similar_group"]].append(q["id"])
for q in questions:
    ids = groups[q["similar_group"]]
    q["similar_count"] = len(ids)
    q["similar_ids"] = [i for i in ids if i != q["id"]]
    q["is_similar"] = len(ids) > 1

# sort hard -> easy inside article, then articles
questions.sort(key=lambda q: (q["article"], DIFF_ORDER[q["difficulty"]], q["id"]))

articles_map = defaultdict(list)
for q in questions:
    articles_map[q["article"]].append(q)

articles = []
for art, qs in articles_map.items():
    qs_sorted = sorted(qs, key=lambda q: (DIFF_ORDER[q["difficulty"]], q["id"]))
    articles.append({
        "key": art,
        "title": qs[0]["article_title"],
        "count": len(qs),
        "hard": sum(1 for x in qs if x["difficulty"]=="سخت"),
        "medium": sum(1 for x in qs if x["difficulty"]=="متوسط"),
        "easy": sum(1 for x in qs if x["difficulty"]=="آسان"),
        "questions": qs_sorted,
    })
# articles with more hard first? user asked questions categorized hard to easy
# list articles by first appearance of hardness then name
articles.sort(key=lambda a: (-a["hard"], -a["medium"], a["key"]))

lesson = {
    "id": "adm",
    "name": "آیین دادرسی مدنی",
    "pack": "بسته ۳",
    "count": len(questions),
    "articles": articles,
}

lessons_index = {
    "lessons": [
        {"id": "madani", "name": "حقوق مدنی", "file": "data/madani.json", "ready": False, "count": 0},
        {"id": "adm", "name": "آیین دادرسی مدنی", "file": "data/adm.json", "ready": True, "count": len(questions)},
        {"id": "tejarat", "name": "حقوق تجارت", "file": "data/tejarat.json", "ready": False, "count": 0},
        {"id": "jaza", "name": "حقوق جزا", "file": "data/jaza.json", "ready": False, "count": 0},
        {"id": "adk", "name": "آیین دادرسی کیفری", "file": "data/adk.json", "ready": False, "count": 0},
        {"id": "osul", "name": "اصول فقه", "file": "data/osul.json", "ready": False, "count": 0},
        {"id": "motun", "name": "متون فقه", "file": "data/motun.json", "ready": False, "count": 0},
    ]
}

os.makedirs(OUT, exist_ok=True)
with open(os.path.join(OUT, "lessons.json"), "w", encoding="utf-8") as f:
    json.dump(lessons_index, f, ensure_ascii=False, indent=2)
with open(os.path.join(OUT, "adm.json"), "w", encoding="utf-8") as f:
    json.dump(lesson, f, ensure_ascii=False, indent=2)

# empty stubs
for lid, name in [("madani","حقوق مدنی"),("tejarat","حقوق تجارت"),("jaza","حقوق جزا"),
                  ("adk","آیین دادرسی کیفری"),("osul","اصول فقه"),("motun","متون فقه")]:
    stub = {"id": lid, "name": name, "pack": "", "count": 0, "articles": []}
    with open(os.path.join(OUT, f"{lid}.json"), "w", encoding="utf-8") as f:
        json.dump(stub, f, ensure_ascii=False, indent=2)

# xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName

wb = Workbook()

# Guide sheet
ws0 = wb.active
ws0.title = "راهنما"
ws0["A1"] = "قالب فایل سوالات — سایت تست وکالت"
ws0["A1"].font = Font(name="Tahoma", size=16, bold=True, color="1B3A4B")
guide = [
"",
"ستون‌های الزامی (از ردیف ۲ به بعد در برگه «سوالات»):",
"id — شناسه یکتا (عدد یا متن)",
"question — متن سوال",
"A / B / C / D — گزینه‌ها",
"answer — حرف گزینه صحیح: A یا B یا C یا D",
"explain — پاسخ تشریحی",
"laws — متن قوانین مرتبط (اختیاری)",
"article — کلید دسته‌بندی، مثلاً: ماده ۱۹ ق.آ.د.م",
"article_title — عنوان کوتاه ماده",
"difficulty — آسان / متوسط / سخت",
"similar_group — شناسه گروه سوالات شبیه هم (اختیاری)",
"",
"برای افزودن درس جدید:",
"۱) از این فایل یک کپی بگیرید با نام درس (مثلا tejarat.xlsx)",
"۲) برگه سوالات را پر کنید",
"۳) در برگه lessons.json سایت، فایل را معرفی کنید",
"۴) یا اسکریپت tools/excel_to_json.py را اجرا کنید",
"",
"ترتیب پیشنهادی سختی:",
"آسان = بازنویسی عین ماده",
"متوسط = ترکیب دو ماده یا نکته مفهوم مخالف",
"سخت = فرض مسئله با اسامی، چند حوزه قضایی، چند قانون",
]
for i, line in enumerate(guide, start=2):
    ws0[f"A{i}"] = line
    ws0[f"A{i}"].font = Font(name="Tahoma", size=11)
ws0.column_dimensions["A"].width = 90
ws0.sheet_view.rightToLeft = True

ws = wb.create_sheet("سوالات")
headers = ["id","lesson","question","A","B","C","D","answer","answer_text","explain","laws",
           "article","article_title","difficulty","similar_group","is_similar","similar_count"]
header_fill = PatternFill("solid", fgColor="1B3A4B")
header_font = Font(name="Tahoma", bold=True, color="FFFFFF", size=10)
thin = Border(
    left=Side(style="thin", color="D0D5DD"),
    right=Side(style="thin", color="D0D5DD"),
    top=Side(style="thin", color="D0D5DD"),
    bottom=Side(style="thin", color="D0D5DD"),
)
# difficulty colors
fills = {
    "سخت": PatternFill("solid", fgColor="FCE8E6"),
    "متوسط": PatternFill("solid", fgColor="FEF3C7"),
    "آسان": PatternFill("solid", fgColor="D1FAE5"),
}
wrap = Alignment(wrap_text=True, vertical="top", horizontal="right")

for col, h in enumerate(headers, 1):
    cell = ws.cell(1, col, h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# sort for excel: article then hard to easy
excel_rows = sorted(questions, key=lambda q: (q["article"], DIFF_ORDER[q["difficulty"]], q["id"]))
for r, q in enumerate(excel_rows, 2):
    vals = [
        q["id"], q["lesson"], q["question"],
        q["options"]["A"], q["options"]["B"], q["options"]["C"], q["options"]["D"],
        q["answer"], q["answer_text"], q["explain"], q["laws"],
        q["article"], q["article_title"], q["difficulty"], q["similar_group"],
        "بله" if q["is_similar"] else "خیر", q["similar_count"],
    ]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(r, c, v)
        cell.alignment = wrap
        cell.font = Font(name="Tahoma", size=10)
        cell.border = thin
        if headers[c-1] == "difficulty":
            cell.fill = fills.get(q["difficulty"], PatternFill())
            cell.alignment = Alignment(horizontal="center", vertical="center")

widths = {
    "A":12,"B":22,"C":55,"D":40,"E":40,"F":40,"G":40,"H":10,"I":40,"J":50,"K":40,
    "L":28,"M":28,"N":12,"O":20,"P":12,"Q":14
}
for col, w in widths.items():
    ws.column_dimensions[col].width = w
ws.row_dimensions[1].height = 24
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:Q{len(excel_rows)+1}"
ws.sheet_view.rightToLeft = True

dv = DataValidation(type="list", formula1='"آسان,متوسط,سخت"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"N2:N{len(excel_rows)+200}")

# summary sheet
ws2 = wb.create_sheet("مواد")
ws2.sheet_view.rightToLeft = True
for col, h in enumerate(["ماده","عنوان","تعداد","سخت","متوسط","آسان"], 1):
    cell = ws2.cell(1, col, h)
    cell.fill = header_fill
    cell.font = header_font
for r, a in enumerate(articles, 2):
    ws2.cell(r,1,a["key"]).font = Font(name="Tahoma")
    ws2.cell(r,2,a["title"]).font = Font(name="Tahoma")
    ws2.cell(r,3,a["count"])
    ws2.cell(r,4,a["hard"])
    ws2.cell(r,5,a["medium"])
    ws2.cell(r,6,a["easy"])
for col in range(1,7):
    ws2.column_dimensions[get_column_letter(col)].width = 28

xlsx_path = os.path.join(OUT, "آیین-دادرسی-مدنی.xlsx")
wb.save(xlsx_path)
print("questions", len(questions))
print("articles", len(articles))
for a in articles:
    print(f"  {a['key']} | {a['title']} | n={a['count']} H{a['hard']} M{a['medium']} E{a['easy']}")
print("saved", xlsx_path)
