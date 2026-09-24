#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Convert a questions xlsx (sheet سوالات) to lesson JSON for the static site."""
import json, sys, os
from collections import defaultdict
from openpyxl import load_workbook

DIFF = {"سخت": 0, "متوسط": 1, "آسان": 2}

def convert(path, lesson_id, lesson_name, out_path):
    wb = load_workbook(path, data_only=True)
    ws = wb["سوالات"] if "سوالات" in wb.sheetnames else wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if not rec.get("question") and not rec.get("سوال"):
            continue
        q = {
            "id": str(rec.get("id") or rec.get("ID") or ""),
            "lesson": lesson_name,
            "question": rec.get("question") or rec.get("سوال"),
            "options": {
                "A": rec.get("A") or rec.get("گزینه A") or "",
                "B": rec.get("B") or rec.get("گزینه B") or "",
                "C": rec.get("C") or rec.get("گزینه C") or "",
                "D": rec.get("D") or rec.get("گزینه D") or "",
            },
            "answer": str(rec.get("answer") or "A").replace("Choice","").strip().upper()[:1],
            "answer_text": rec.get("answer_text") or rec.get("پاسخ صحیح") or "",
            "explain": rec.get("explain") or rec.get("پاسخ تشریحی") or "",
            "laws": rec.get("laws") or rec.get("قوانین مرتبط") or "",
            "article": rec.get("article") or "سایر",
            "article_title": rec.get("article_title") or "",
            "difficulty": rec.get("difficulty") or "متوسط",
            "similar_group": rec.get("similar_group") or "",
            "is_similar": str(rec.get("is_similar") or "") == "بله",
            "similar_count": int(rec.get("similar_count") or 1),
        }
        questions.append(q)
    amap = defaultdict(list)
    for q in questions:
        amap[q["article"]].append(q)
    articles = []
    for key, qs in amap.items():
        qs = sorted(qs, key=lambda x: DIFF.get(x["difficulty"], 1))
        articles.append({
            "key": key,
            "title": qs[0]["article_title"] or key,
            "count": len(qs),
            "hard": sum(1 for x in qs if x["difficulty"]=="سخت"),
            "medium": sum(1 for x in qs if x["difficulty"]=="متوسط"),
            "easy": sum(1 for x in qs if x["difficulty"]=="آسان"),
            "questions": qs,
        })
    articles.sort(key=lambda a: (-a["hard"], -a["medium"], a["key"]))
    data = {"id": lesson_id, "name": lesson_name, "pack": "", "count": len(questions), "articles": articles}
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print("wrote", out_path, "questions=", len(questions), "articles=", len(articles))

if __name__ == "__main__":
    if len(sys.argv) < 5:
        print("usage: excel_to_json.py input.xlsx lesson_id 'نام درس' out.json")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4])
